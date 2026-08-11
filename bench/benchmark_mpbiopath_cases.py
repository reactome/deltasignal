#!/usr/bin/env python3
"""Benchmark DeltaSignal against fixed MP-BioPath perturbation cases.

The benchmark keeps the network catalog fixed while solver configurations or
DeltaSignal commits change. It joins the published supplementary case table to
LNG UUIDs through Reactome database identifiers and reports all mapping losses.
Missing mappings are never converted into synthetic NORMAL predictions.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import random
import re
import signal
import subprocess
import sys
import time
from collections import Counter, defaultdict, deque
from dataclasses import asdict, dataclass
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen


DOWN, NORMAL, UP = 0, 1, 2
LABELS = (DOWN, NORMAL, UP)
STATUS_TO_CLASS = {
    "downregulation/inactivation": DOWN,
    "unchanged": NORMAL,
    "upregulation/activation": UP,
}
DEVELOPMENT_PATHWAYS = (
    "R-HSA-1257604",  # PIP3 activates AKT signaling
    "R-HSA-453279",   # Mitotic G1 phase and G1/S transition
    "R-HSA-69620",    # Cell Cycle Checkpoints
)
HELD_OUT_PATHWAYS = (
    "R-HSA-68875",    # Mitotic Prophase
    "R-HSA-69242",    # S Phase
    "R-HSA-195721",   # Signaling by WNT
    "R-HSA-1227986",  # Signaling by ERBB2
    "R-HSA-3700989",  # Transcriptional Regulation by TP53
    "R-HSA-5673001",  # RAF/MAP kinase cascade
    "R-HSA-5693567",  # HDR through HRR or SSA
)
DEFAULT_PATHWAYS = DEVELOPMENT_PATHWAYS + HELD_OUT_PATHWAYS


@dataclass(frozen=True)
class Case:
    pathway_id: str
    pathway_name: str
    gene: str
    direction: int
    key_output_dbid: str
    expected: int
    curator_prediction: int
    mpbiopath_prediction: int
    evidence: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--supplementary-workbook", type=Path, required=True)
    parser.add_argument("--id-map", type=Path, required=True)
    parser.add_argument("--catalog", type=Path, required=True)
    parser.add_argument(
        "--reactome-id-audit",
        type=Path,
        help="Optional TSV with dbid, stable_id, schema_class, and display_name from the pinned Reactome release",
    )
    parser.add_argument("--deltasignal-dir", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--pathway-id", action="append", dest="pathway_ids")
    parser.add_argument("--ground-truth", choices=("experimental", "curator"), default="experimental")
    parser.add_argument("--down-cutoff", type=float, default=0.85)
    parser.add_argument("--up-cutoff", type=float, default=1.15)
    parser.add_argument("--output-aggregation", choices=("max", "mean", "min", "extreme"), default="max")
    parser.add_argument(
        "--allow-output-proxies",
        action="store_true",
        help="Score missing key-output entities through LNG's explicit entity-to-reaction proxy export",
    )
    parser.add_argument(
        "--prefer-output-proxies",
        action="store_true",
        help="Audit proxy validity by using an available reaction proxy even when the exact entity is exported",
    )
    parser.add_argument(
        "--derive-output-proxies",
        action="store_true",
        help="Derive producing/consuming reaction proxies from graph adjacency for exact-entity proxy validation",
    )
    parser.add_argument("--perturbation-up", type=float, default=80.0)
    parser.add_argument("--port", type=int, default=18080)
    parser.add_argument("--bootstrap", type=int, default=2000)
    parser.add_argument("--seed", type=int, default=20260721)
    parser.add_argument("--server-timeout", type=float, default=180.0)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def git_state(path: Path) -> dict[str, object]:
    def run(*args: str) -> str:
        return subprocess.check_output(
            ["git", "-C", str(path), *args], text=True, stderr=subprocess.DEVNULL
        ).strip()

    try:
        status = run("status", "--porcelain").splitlines()
        return {
            "commit": run("rev-parse", "HEAD"),
            "branch": run("branch", "--show-current"),
            "dirty": bool(status),
            "status_porcelain": status,
        }
    except (OSError, subprocess.CalledProcessError):
        return {"available": False}


def normalize_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_cases(workbook: Path, ground_truth: str) -> list[Case]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the supplementary workbook") from exc

    sheet = load_workbook(workbook, read_only=True, data_only=True)["Supplementary Table S1"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    index = {str(name): idx for idx, name in enumerate(header)}
    required = {
        "Reactome_pathway_ID",
        "Reactome_pathway_name",
        "RootInput_Perturbation",
        "KeyOutput_ID",
        "CuratorPredictedKeyOutputStatus",
        "MP-BioPathPredictedKeyOutputStatus",
        "ExperimentalKeyOutputStatus",
        "PublishedEvidence",
    }
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(f"Supplementary Table S1 is missing columns: {missing}")

    cases: list[Case] = []
    truth_column = (
        "ExperimentalKeyOutputStatus"
        if ground_truth == "experimental"
        else "CuratorPredictedKeyOutputStatus"
    )
    perturbation_pattern = re.compile(r"^(.+)_(downregulation|upregulation)$")
    for row in rows:
        expected_text = normalize_cell(row[index[truth_column]])
        if expected_text not in STATUS_TO_CLASS:
            continue
        perturbation = normalize_cell(row[index["RootInput_Perturbation"]])
        match = perturbation_pattern.fullmatch(perturbation)
        if match is None:
            raise ValueError(f"Unexpected perturbation label: {perturbation!r}")
        gene, direction_text = match.groups()
        curator_text = normalize_cell(row[index["CuratorPredictedKeyOutputStatus"]])
        mpbiopath_text = normalize_cell(row[index["MP-BioPathPredictedKeyOutputStatus"]])
        if curator_text not in STATUS_TO_CLASS or mpbiopath_text not in STATUS_TO_CLASS:
            raise ValueError(f"Unexpected prediction status in case {perturbation}")
        pathway_numeric = normalize_cell(row[index["Reactome_pathway_ID"]])
        cases.append(
            Case(
                pathway_id=f"R-HSA-{pathway_numeric}",
                pathway_name=normalize_cell(row[index["Reactome_pathway_name"]]),
                gene=gene,
                direction=DOWN if direction_text == "downregulation" else UP,
                key_output_dbid=normalize_cell(row[index["KeyOutput_ID"]]),
                expected=STATUS_TO_CLASS[expected_text],
                curator_prediction=STATUS_TO_CLASS[curator_text],
                mpbiopath_prediction=STATUS_TO_CLASS[mpbiopath_text],
                evidence=normalize_cell(row[index["PublishedEvidence"]]),
            )
        )
    return cases


def load_gene_dbids(id_map: Path, genes: set[str]) -> dict[str, set[str]]:
    mapped: dict[str, set[str]] = defaultdict(set)
    with id_map.open(newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"Database_Identifier", "Display_Name", "Reference_Entity_Name"}
        missing = sorted(required - set(reader.fieldnames or []))
        if missing:
            raise ValueError(f"ID map is missing columns: {missing}")
        for row in reader:
            names = {normalize_cell(row.get("Display_Name")), normalize_cell(row.get("Reference_Entity_Name"))}
            for gene in names & genes:
                mapped[gene].add(normalize_cell(row["Database_Identifier"]))
    return mapped


def find_pathway_dir(catalog: Path, pathway_id: str) -> Path | None:
    matches = sorted(path for path in catalog.iterdir() if path.is_dir() and path.name.endswith(pathway_id))
    if len(matches) > 1:
        raise ValueError(f"Multiple catalog directories match {pathway_id}: {matches}")
    return matches[0] if matches else None


def load_dbid_to_uuids(pathway_dir: Path) -> dict[str, list[str]]:
    mapping: dict[str, list[str]] = defaultdict(list)
    seen: set[tuple[str, str]] = set()
    nodes_file = pathway_dir / "nodes.csv"
    if nodes_file.exists():
        with nodes_file.open(newline="") as handle:
            for row in csv.DictReader(handle):
                stable_ids = {normalize_cell(row.get("diagram_entity_id"))}
                stable_ids.update(normalize_cell(item) for item in normalize_cell(row.get("member_leaves")).split("|"))
                for stable_id in stable_ids - {""}:
                    dbid = stable_id.rsplit("-", 1)[-1]
                    pair = (dbid, row["uuid"])
                    if pair not in seen:
                        seen.add(pair)
                        mapping[dbid].append(row["uuid"])
    else:
        with (pathway_dir / "stid_to_uuid_mapping.csv").open(newline="") as handle:
            for row in csv.DictReader(handle):
                dbid = row["stable_id"].rsplit("-", 1)[-1]
                pair = (dbid, row["uuid"])
                if pair not in seen:
                    seen.add(pair)
                    mapping[dbid].append(row["uuid"])
    return mapping


def load_proxy_dbid_to_uuids(pathway_dir: Path) -> dict[str, dict[str, list[str]]]:
    """Load explicit LNG entity-to-reaction proxies, grouped by proxy role."""
    proxy_file = pathway_dir / "entity_reaction_proxy_mapping.csv"
    mapping: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    if not proxy_file.exists():
        return mapping
    with proxy_file.open(newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"entity_stable_id", "proxy_uuid", "proxy_role"}
        missing = sorted(required - set(reader.fieldnames or []))
        if missing:
            raise ValueError(f"Proxy map is missing columns: {missing}")
        for row in reader:
            dbid = normalize_cell(row["entity_stable_id"]).rsplit("-", 1)[-1]
            role = normalize_cell(row["proxy_role"])
            uuid = normalize_cell(row["proxy_uuid"])
            if dbid and role and uuid and uuid not in mapping[dbid][role]:
                mapping[dbid][role].append(uuid)
    return mapping


def derive_proxy_dbid_to_uuids(pathway_dir: Path) -> dict[str, dict[str, list[str]]]:
    """Derive adjacent reaction proxies for entities already present in a graph."""
    dbid_to_uuids = load_dbid_to_uuids(pathway_dir)
    node_kinds: dict[str, str] = {}
    with (pathway_dir / "nodes.csv").open(newline="") as handle:
        for row in csv.DictReader(handle):
            node_kinds[normalize_cell(row["uuid"])] = normalize_cell(row["node_kind"])
    entity_to_dbids: dict[str, set[str]] = defaultdict(set)
    for dbid, uuids in dbid_to_uuids.items():
        for uuid in uuids:
            if node_kinds.get(uuid) != "reaction":
                entity_to_dbids[uuid].add(dbid)
    mapping: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    with (pathway_dir / "logic_network.csv").open(newline="") as handle:
        for row in csv.DictReader(handle):
            source = normalize_cell(row["source_id"])
            target = normalize_cell(row["target_id"])
            if node_kinds.get(source) == "reaction" and target in entity_to_dbids:
                for dbid in entity_to_dbids[target]:
                    if source not in mapping[dbid]["producing"]:
                        mapping[dbid]["producing"].append(source)
            if source in entity_to_dbids and node_kinds.get(target) == "reaction":
                for dbid in entity_to_dbids[source]:
                    if target not in mapping[dbid]["consuming"]:
                        mapping[dbid]["consuming"].append(target)
    return mapping


def load_reactome_id_audit(path: Path | None) -> dict[str, dict[str, str]]:
    """Load identifiers exported from the exact Reactome release used by LNG."""
    if path is None:
        return {}
    audit: dict[str, dict[str, str]] = {}
    with path.open(newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"dbid", "stable_id", "schema_class", "display_name"}
        missing = sorted(required - set(reader.fieldnames or []))
        if missing:
            raise ValueError(f"Reactome ID audit is missing columns: {missing}")
        for row in reader:
            dbid = normalize_cell(row["dbid"])
            if dbid:
                audit[dbid] = {key: normalize_cell(value) for key, value in row.items()}
    return audit


def load_network_adjacency(pathway_dir: Path) -> dict[str, list[tuple[str, int]]]:
    """Load the signed directed graph used by the structural baselines."""
    adjacency: dict[str, list[tuple[str, int]]] = defaultdict(list)
    with (pathway_dir / "logic_network.csv").open(newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"source_id", "target_id", "pos_neg"}
        missing = sorted(required - set(reader.fieldnames or []))
        if missing:
            raise ValueError(f"Logic network is missing columns: {missing}")
        for row in reader:
            sign = -1 if normalize_cell(row["pos_neg"]) == "neg" else 1
            adjacency[normalize_cell(row["source_id"])].append(
                (normalize_cell(row["target_id"]), sign)
            )
    return adjacency


def reachable_path_signs(
    adjacency: dict[str, list[tuple[str, int]]],
    sources: list[str],
    targets: list[str],
    *,
    shortest_only: bool,
) -> set[int]:
    """Return signed path products without enumerating cyclic paths."""
    target_set = set(targets)
    queue = deque((source, 1, 0) for source in sources)
    seen_distance: dict[tuple[str, int], int] = {(source, 1): 0 for source in sources}
    found: list[tuple[int, int]] = []
    while queue:
        node, path_sign, distance = queue.popleft()
        if node in target_set:
            found.append((distance, path_sign))
        for target, edge_sign in adjacency.get(node, ()):
            state = (target, path_sign * edge_sign)
            next_distance = distance + 1
            if state in seen_distance and seen_distance[state] <= next_distance:
                continue
            seen_distance[state] = next_distance
            queue.append((target, state[1], next_distance))
    if not found:
        return set()
    if shortest_only:
        minimum = min(distance for distance, _ in found)
        return {sign for distance, sign in found if distance == minimum}
    return {sign for _, sign in found}


def structural_prediction(path_signs: set[int], perturbation_direction: int) -> int:
    """Convert path polarity into a conservative three-class prediction."""
    if len(path_signs) != 1:
        return NORMAL
    path_sign = next(iter(path_signs))
    perturbation_sign = -1 if perturbation_direction == DOWN else 1
    return UP if path_sign * perturbation_sign > 0 else DOWN


def request_json(url: str, payload: dict | None = None, timeout: float = 600.0) -> dict:
    data = None if payload is None else json.dumps(payload).encode()
    request = Request(url, data=data, headers={"Content-Type": "application/json"})
    with urlopen(request, timeout=timeout) as response:
        return json.loads(response.read())


def start_server(args: argparse.Namespace) -> tuple[subprocess.Popen, Path]:
    args.output_dir.mkdir(parents=True, exist_ok=True)
    log_path = args.output_dir / "deltasignal_server.log"
    log_handle = log_path.open("w")
    env = os.environ.copy()
    env["DS_PATHWAY_CATALOG"] = str(args.catalog.resolve())
    # Display-name enrichment is unrelated to numerical predictions and makes
    # an otherwise local benchmark depend on the Reactome web service.
    env["DS_REACTOME_ENRICH"] = "0"
    process = subprocess.Popen(
        [
            "julia",
            f"--project={args.deltasignal_dir.resolve()}",
            str(args.deltasignal_dir.resolve() / "src/api/server.jl"),
            "--port",
            str(args.port),
        ],
        cwd=args.deltasignal_dir,
        env=env,
        stdout=log_handle,
        stderr=subprocess.STDOUT,
        text=True,
    )
    process._deltasignal_log_handle = log_handle  # type: ignore[attr-defined]
    base = f"http://127.0.0.1:{args.port}"
    deadline = time.time() + args.server_timeout
    while time.time() < deadline:
        if process.poll() is not None:
            log_handle.flush()
            raise RuntimeError(f"DeltaSignal server exited early; see {log_path}")
        try:
            if request_json(f"{base}/api/health", timeout=2).get("status") == "ok":
                return process, log_path
        except (HTTPError, URLError, TimeoutError):
            time.sleep(0.5)
    process.terminate()
    raise TimeoutError(f"DeltaSignal server did not become ready; see {log_path}")


def stop_server(process: subprocess.Popen) -> None:
    if process.poll() is None:
        process.send_signal(signal.SIGINT)
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.terminate()
            process.wait(timeout=10)
    log_handle = getattr(process, "_deltasignal_log_handle", None)
    if log_handle is not None:
        log_handle.close()


def aggregate(values: list[float], mode: str) -> float:
    if mode == "mean":
        return sum(values) / len(values)
    if mode == "min":
        return min(values)
    if mode == "extreme":
        return max(values, key=lambda value: abs(value - 1.0))
    return max(values)


def classify(value: float, down_cutoff: float, up_cutoff: float) -> int:
    if value < down_cutoff:
        return DOWN
    if value >= up_cutoff:
        return UP
    return NORMAL


def metric_summary(rows: list[dict[str, object]]) -> dict[str, object]:
    scored = [row for row in rows if row["prediction"] is not None]
    converged = [row for row in scored if bool(row["converged"])]
    confusion = Counter((int(row["prediction"]), int(row["expected"])) for row in scored)
    per_class = {}
    for label in LABELS:
        tp = confusion[(label, label)]
        fp = sum(confusion[(label, expected)] for expected in LABELS if expected != label)
        fn = sum(confusion[(predicted, label)] for predicted in LABELS if predicted != label)
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
        per_class[str(label)] = {"precision": precision, "recall": recall, "f1": f1, "support": tp + fn}
    correct = sum(int(row["prediction"]) == int(row["expected"]) for row in scored)
    converged_correct = sum(
        int(row["prediction"]) == int(row["expected"]) for row in converged
    )
    solve_states = {
        (str(row["pathway_id"]), str(row["gene"]), int(row["direction"])): bool(
            row["converged"]
        )
        for row in scored
        if all(key in row for key in ("pathway_id", "gene", "direction"))
    }
    return {
        "eligible_cases": len(rows),
        "scored_cases": len(scored),
        "coverage": len(scored) / len(rows) if rows else 0.0,
        "correct": correct,
        "accuracy": correct / len(scored) if scored else 0.0,
        "coverage_adjusted_accuracy": correct / len(rows) if rows else 0.0,
        "converged_scored_cases": len(converged),
        "converged_correct": converged_correct,
        "converged_accuracy": (
            converged_correct / len(converged) if converged else 0.0
        ),
        "converged_coverage": len(converged) / len(rows) if rows else 0.0,
        "macro_f1": sum(per_class[str(label)]["f1"] for label in LABELS) / len(LABELS),
        "balanced_accuracy": sum(per_class[str(label)]["recall"] for label in LABELS) / len(LABELS),
        "change_f1": (per_class[str(DOWN)]["f1"] + per_class[str(UP)]["f1"]) / 2,
        "per_class": per_class,
        "confusion": {f"pred_{predicted}_expected_{expected}": count for (predicted, expected), count in sorted(confusion.items())},
        "unscored_reasons": dict(Counter(str(row["mapping_status"]) for row in rows if row["prediction"] is None)),
        "nonconverged_cases": len(scored) - len(converged),
        # Retained for compatibility with earlier result readers. This is a
        # case count; unique perturbation solves are reported separately.
        "nonconverged_solves": len(scored) - len(converged),
        "unique_scored_solves": len(solve_states),
        "unique_nonconverged_solves": sum(
            not converged for converged in solve_states.values()
        ),
    }


def metric_summary_for_field(
    rows: list[dict[str, object]],
    field: str,
) -> dict[str, object]:
    projected = [
        {
            **row,
            "prediction": row.get(field),
            "mapping_status": (
                row.get("mapping_status", "mapped")
                if row.get(field) is None
                else "scored"
            ),
            "converged": True,
        }
        for row in rows
    ]
    return metric_summary(projected)


def grouped_summaries(
    rows: list[dict[str, object]],
    field: str = "prediction",
) -> dict[str, dict[str, object]]:
    groups: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[str(row["pathway_id"])].append(row)
    return {
        pathway_id: metric_summary_for_field(group_rows, field)
        for pathway_id, group_rows in sorted(groups.items())
    }


def split_summaries(
    rows: list[dict[str, object]],
    field: str = "prediction",
) -> dict[str, dict[str, object]]:
    development = [
        row for row in rows if str(row["pathway_id"]) in DEVELOPMENT_PATHWAYS
    ]
    held_out = [
        row for row in rows if str(row["pathway_id"]) in HELD_OUT_PATHWAYS
    ]
    return {
        "development": metric_summary_for_field(development, field),
        "held_out": metric_summary_for_field(held_out, field),
    }


def baseline_accuracy(
    rows: list[dict[str, object]],
    field: str,
    *,
    deltasignal_scored_only: bool = True,
) -> dict[str, float | int]:
    """Score a published baseline on either the paired or full eligible set."""
    evaluated = (
        [row for row in rows if row["prediction"] is not None]
        if deltasignal_scored_only
        else rows
    )
    correct = sum(int(row[field]) == int(row["expected"]) for row in evaluated)
    return {
        "correct": correct,
        "total": len(evaluated),
        "accuracy": correct / len(evaluated) if evaluated else 0.0,
    }


def paired_bootstrap(rows: list[dict[str, object]], field: str, iterations: int, seed: int) -> dict[str, float]:
    scored = [row for row in rows if row["prediction"] is not None]
    if not scored or iterations <= 0:
        return {}
    rng = random.Random(seed)
    differences = []
    for _ in range(iterations):
        sample = [scored[rng.randrange(len(scored))] for _ in scored]
        ds = sum(int(row["prediction"]) == int(row["expected"]) for row in sample) / len(sample)
        baseline = sum(int(row[field]) == int(row["expected"]) for row in sample) / len(sample)
        differences.append(ds - baseline)
    differences.sort()
    return {
        "mean_accuracy_difference": sum(differences) / len(differences),
        "ci95_low": differences[math.floor(0.025 * (len(differences) - 1))],
        "ci95_high": differences[math.ceil(0.975 * (len(differences) - 1))],
    }


def run(args: argparse.Namespace) -> dict[str, object]:
    requested = set(args.pathway_ids or DEFAULT_PATHWAYS)
    all_cases = load_cases(args.supplementary_workbook, args.ground_truth)
    cases = [case for case in all_cases if case.pathway_id in requested]
    found_ids = {case.pathway_id for case in cases}
    if requested - found_ids:
        raise ValueError(f"Requested pathways absent from benchmark cases: {sorted(requested - found_ids)}")

    gene_dbids = load_gene_dbids(args.id_map, {case.gene for case in cases})
    reactome_id_audit = load_reactome_id_audit(args.reactome_id_audit)
    pathway_dirs = {pathway_id: find_pathway_dir(args.catalog, pathway_id) for pathway_id in requested}
    absent_networks = sorted(pathway_id for pathway_id, path in pathway_dirs.items() if path is None)
    if absent_networks:
        raise FileNotFoundError(f"No LNG network for pathways: {absent_networks}")
    dbid_maps = {pathway_id: load_dbid_to_uuids(path) for pathway_id, path in pathway_dirs.items() if path}
    proxy_loader = (
        derive_proxy_dbid_to_uuids
        if args.derive_output_proxies
        else load_proxy_dbid_to_uuids
    )
    proxy_maps = {
        pathway_id: proxy_loader(path)
        for pathway_id, path in pathway_dirs.items()
        if path
    }
    adjacencies = {
        pathway_id: load_network_adjacency(path)
        for pathway_id, path in pathway_dirs.items()
        if path
    }

    process, log_path = start_server(args)
    base = f"http://127.0.0.1:{args.port}"
    parsed_ids = {}
    rows: list[dict[str, object]] = []
    solve_cache: dict[tuple[str, str, int], dict] = {}
    try:
        for pathway_id, pathway_dir in pathway_dirs.items():
            parsed = request_json(f"{base}/api/parse", {"pathway_id": pathway_dir.name})
            if parsed.get("status") != "success":
                raise RuntimeError(f"Parse failed for {pathway_id}: {parsed}")
            parsed_ids[pathway_id] = parsed["network_id"]

        for case in cases:
            dbid_map = dbid_maps[case.pathway_id]
            gene_uuids = sorted({uuid for dbid in gene_dbids.get(case.gene, ()) for uuid in dbid_map.get(dbid, ())})
            exact_output_uuids = sorted(set(dbid_map.get(case.key_output_dbid, ())))
            proxy_roles = proxy_maps[case.pathway_id].get(case.key_output_dbid, {})
            preferred_proxy_role = next(
                (
                    role
                    for role in ("producing", "consuming")
                    if proxy_roles.get(role)
                ),
                next(iter(sorted(proxy_roles)), None),
            )
            proxy_output_uuids = (
                sorted(set(proxy_roles[preferred_proxy_role]))
                if preferred_proxy_role is not None
                else []
            )
            if proxy_output_uuids and args.prefer_output_proxies:
                output_uuids = proxy_output_uuids
                output_mapping_mode = f"proxy_{preferred_proxy_role}_preferred"
            elif exact_output_uuids:
                output_uuids = exact_output_uuids
                output_mapping_mode = "exact"
            elif proxy_output_uuids and args.allow_output_proxies:
                output_uuids = proxy_output_uuids
                output_mapping_mode = f"proxy_{preferred_proxy_role}"
            elif proxy_output_uuids:
                output_uuids = []
                output_mapping_mode = "proxy_available_not_enabled"
            else:
                output_uuids = []
                output_mapping_mode = "absent_from_network"
            all_path_signs = (
                reachable_path_signs(
                    adjacencies[case.pathway_id],
                    gene_uuids,
                    output_uuids,
                    shortest_only=False,
                )
                if gene_uuids and output_uuids
                else set()
            )
            shortest_path_signs = (
                reachable_path_signs(
                    adjacencies[case.pathway_id],
                    gene_uuids,
                    output_uuids,
                    shortest_only=True,
                )
                if gene_uuids and output_uuids
                else set()
            )
            reactome_output = reactome_id_audit.get(case.key_output_dbid, {})
            row = asdict(case)
            row.update(
                {
                    "evaluation_split": (
                        "development"
                        if case.pathway_id in DEVELOPMENT_PATHWAYS
                        else "held_out"
                    ),
                    "gene_dbid_count": len(gene_dbids.get(case.gene, ())),
                    "gene_uuid_count": len(gene_uuids),
                    "exact_output_uuid_count": len(exact_output_uuids),
                    "proxy_output_uuid_count": len(proxy_output_uuids),
                    "output_uuid_count": len(output_uuids),
                    "output_mapping_mode": output_mapping_mode,
                    "reactome_output_present": (
                        case.key_output_dbid in reactome_id_audit
                        if reactome_id_audit
                        else None
                    ),
                    "reactome_output_stable_id": reactome_output.get("stable_id"),
                    "reactome_output_schema_class": reactome_output.get("schema_class"),
                    "connectivity_status": (
                        "directed_path"
                        if all_path_signs
                        else "no_directed_path"
                        if gene_uuids and output_uuids
                        else "not_testable"
                    ),
                    "reachable_path_signs": ",".join(str(sign) for sign in sorted(all_path_signs)),
                    "shortest_path_signs": ",".join(
                        str(sign) for sign in sorted(shortest_path_signs)
                    ),
                    "no_change_prediction": NORMAL,
                    "class_frequency_prediction": None,
                    "signed_reachability_prediction": (
                        structural_prediction(all_path_signs, case.direction)
                        if gene_uuids and output_uuids
                        else None
                    ),
                    "shortest_signed_path_prediction": (
                        structural_prediction(shortest_path_signs, case.direction)
                        if gene_uuids and output_uuids
                        else None
                    ),
                    "prediction": None,
                    "predicted_ui": None,
                    "converged": None,
                    "iterations": None,
                    "mapping_status": "mapped",
                }
            )
            if not gene_dbids.get(case.gene):
                row["mapping_status"] = "gene_absent_from_legacy_id_map"
            elif not gene_uuids:
                if reactome_id_audit and not (
                    set(gene_dbids[case.gene]) & set(reactome_id_audit)
                ):
                    row["mapping_status"] = "gene_dbids_absent_from_reactome_release"
                else:
                    row["mapping_status"] = "gene_dbids_present_in_reactome_but_not_exported"
            elif not output_uuids:
                if proxy_output_uuids:
                    row["mapping_status"] = "key_output_proxy_available_not_enabled"
                elif reactome_id_audit and case.key_output_dbid not in reactome_id_audit:
                    row["mapping_status"] = "key_output_absent_from_reactome_release"
                elif reactome_id_audit:
                    row["mapping_status"] = "key_output_present_in_reactome_but_not_exported"
                else:
                    row["mapping_status"] = "key_output_dbid_absent_from_network"
            else:
                cache_key = (case.pathway_id, case.gene, case.direction)
                if cache_key not in solve_cache:
                    perturbation = 0.0 if case.direction == DOWN else args.perturbation_up
                    observations = {uuid: [perturbation, 1.0] for uuid in gene_uuids}
                    solve_cache[cache_key] = request_json(
                        f"{base}/api/solve",
                        {"network_id": parsed_ids[case.pathway_id], "observations": observations},
                    )
                result = solve_cache[cache_key]
                if result.get("status") != "success":
                    row["mapping_status"] = "solve_failed"
                else:
                    missing_outputs = [
                        uuid
                        for uuid in output_uuids
                        if uuid not in result["node_activities"]
                    ]
                    if missing_outputs:
                        row["mapping_status"] = "solver_output_uuid_missing"
                    else:
                        values = [
                            float(result["node_activities"][uuid]) * 100.0
                            for uuid in output_uuids
                        ]
                        predicted_ui = aggregate(values, args.output_aggregation)
                        row.update(
                            {
                                "prediction": classify(
                                    predicted_ui,
                                    args.down_cutoff,
                                    args.up_cutoff,
                                ),
                                "predicted_ui": predicted_ui,
                                "converged": bool(result.get("converged")),
                                "iterations": result.get("iterations"),
                            }
                        )
            rows.append(row)
    finally:
        stop_server(process)

    development_counts = Counter(
        int(row["expected"])
        for row in rows
        if row["evaluation_split"] == "development"
    )
    class_frequency_prediction = max(
        LABELS,
        key=lambda label: (development_counts[label], label == NORMAL, -label),
    )
    for row in rows:
        row["class_frequency_prediction"] = class_frequency_prediction

    summary = metric_summary(rows)
    summary["evaluation_split"] = {
        "development_pathways": list(DEVELOPMENT_PATHWAYS),
        "held_out_pathways": list(HELD_OUT_PATHWAYS),
        "held_out_was_used_for_configuration": False,
    }
    summary["by_split"] = split_summaries(rows)
    summary["by_pathway"] = grouped_summaries(rows)
    summary["output_mapping_modes"] = dict(
        Counter(str(row["output_mapping_mode"]) for row in rows)
    )
    summary["connectivity_statuses"] = dict(
        Counter(str(row["connectivity_status"]) for row in rows)
    )
    summary["baselines"] = {
        "no_change": {
            "all": metric_summary_for_field(rows, "no_change_prediction"),
            "by_split": split_summaries(rows, "no_change_prediction"),
        },
        "development_class_frequency": {
            "prediction": class_frequency_prediction,
            "development_class_counts": {
                str(label): development_counts[label] for label in LABELS
            },
            "all": metric_summary_for_field(rows, "class_frequency_prediction"),
            "by_split": split_summaries(rows, "class_frequency_prediction"),
        },
        "signed_reachability": {
            "all": metric_summary_for_field(
                rows,
                "signed_reachability_prediction",
            ),
            "by_split": split_summaries(
                rows,
                "signed_reachability_prediction",
            ),
        },
        "shortest_signed_path": {
            "all": metric_summary_for_field(
                rows,
                "shortest_signed_path_prediction",
            ),
            "by_split": split_summaries(
                rows,
                "shortest_signed_path_prediction",
            ),
        },
    }
    summary["curator_on_scored_cases"] = baseline_accuracy(rows, "curator_prediction")
    summary["mpbiopath_on_scored_cases"] = baseline_accuracy(rows, "mpbiopath_prediction")
    summary["curator_on_all_eligible_cases"] = baseline_accuracy(
        rows,
        "curator_prediction",
        deltasignal_scored_only=False,
    )
    summary["mpbiopath_on_all_eligible_cases"] = baseline_accuracy(
        rows,
        "mpbiopath_prediction",
        deltasignal_scored_only=False,
    )
    summary["paired_bootstrap_vs_curator"] = paired_bootstrap(rows, "curator_prediction", args.bootstrap, args.seed)
    summary["paired_bootstrap_vs_mpbiopath"] = paired_bootstrap(rows, "mpbiopath_prediction", args.bootstrap, args.seed + 1)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    case_path = args.output_dir / "benchmark_cases.tsv"
    with case_path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)
    (args.output_dir / "benchmark_summary.json").write_text(json.dumps(summary, indent=2) + "\n")

    network_files = {}
    for pathway_id, pathway_dir in pathway_dirs.items():
        network_files[pathway_id] = {
            name: {"path": str(path), "sha256": sha256(path), "size_bytes": path.stat().st_size}
            for name in (
                "logic_network.csv",
                "stid_to_uuid_mapping.csv",
                "entity_reaction_proxy_mapping.csv",
                "nodes.csv",
            )
            if (path := pathway_dir / name).exists()
        }
    solver_config = {
        key: value for key, value in sorted(os.environ.items()) if key.startswith("DS_")
    }
    solver_config["DS_PATHWAY_CATALOG"] = str(args.catalog.resolve())
    manifest = {
        "created_unix": time.time(),
        "ground_truth": args.ground_truth,
        "pathway_ids": sorted(requested),
        "evaluation_split": {
            "development_pathways": list(DEVELOPMENT_PATHWAYS),
            "held_out_pathways": list(HELD_OUT_PATHWAYS),
        },
        "thresholds": {"down": args.down_cutoff, "up": args.up_cutoff},
        "output_aggregation": args.output_aggregation,
        "allow_output_proxies": args.allow_output_proxies,
        "prefer_output_proxies": args.prefer_output_proxies,
        "derive_output_proxies": args.derive_output_proxies,
        "perturbation_up": args.perturbation_up,
        "solver_environment_overrides": solver_config,
        "deltasignal_git": git_state(args.deltasignal_dir),
        "benchmark_code": {
            "path": str(Path(__file__).resolve()),
            "sha256": sha256(Path(__file__).resolve()),
            "python": sys.version,
        },
        "inputs": {
            "supplementary_workbook": {"path": str(args.supplementary_workbook), "sha256": sha256(args.supplementary_workbook)},
            "id_map": {"path": str(args.id_map), "sha256": sha256(args.id_map)},
            "reactome_id_audit": (
                {
                    "path": str(args.reactome_id_audit),
                    "sha256": sha256(args.reactome_id_audit),
                }
                if args.reactome_id_audit is not None
                else None
            ),
            "networks": network_files,
        },
        "server_log": str(log_path),
        "outputs": {"cases": str(case_path), "summary": str(args.output_dir / "benchmark_summary.json")},
    }
    (args.output_dir / "benchmark_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return summary


def main() -> None:
    args = parse_args()
    summary = run(args)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
